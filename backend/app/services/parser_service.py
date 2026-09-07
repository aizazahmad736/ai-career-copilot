import io
import re
from typing import Tuple
from pypdf import PdfReader
from docx import Document

class DocumentParserService:
    @staticmethod
    def clean_text(text: str) -> str:
        if not text:
            return ""
        # Remove null characters and non-printable control characters except whitespace
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)
        # Fix hyphenated words broken across linebreaks: e.g. "experi-\nence" -> "experience"
        text = re.sub(r'(\w+)-\n(\w+)', r'\1\2', text)
        # Normalize multiple horizontal whitespaces
        text = re.sub(r'[ \t]+', ' ', text)
        # Normalize excessive newlines
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> str:
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text.strip())
            combined = "\n\n".join(pages_text)
            return cls.clean_text(combined)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> str:
        try:
            doc = Document(io.BytesIO(file_bytes))
            lines = []
            for p in doc.paragraphs:
                if p.text.strip():
                    lines.append(p.text.strip())
            
            # Also extract text from any tables
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        lines.append(" | ".join(row_data))

            return cls.clean_text("\n".join(lines))
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

    @classmethod
    def parse_txt(cls, file_bytes: bytes) -> str:
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                decoded = file_bytes.decode(encoding)
                return cls.clean_text(decoded)
            except UnicodeDecodeError:
                continue
        raise ValueError("Failed to decode text file. Ensure it is a valid text format.")

    @classmethod
    def parse_file(cls, file_bytes: bytes, filename: str) -> Tuple[str, str]:
        ext = filename.lower().split(".")[-1] if "." in filename else ""
        if ext == "pdf":
            text = cls.parse_pdf(file_bytes)
            file_type = "pdf"
        elif ext in ["docx", "doc"]:
            text = cls.parse_docx(file_bytes)
            file_type = "docx"
        elif ext in ["txt", "md", "rtf"]:
            text = cls.parse_txt(file_bytes)
            file_type = "txt"
        else:
            raise ValueError(f"Unsupported file format '.{ext}'. Please upload a PDF, DOCX, or TXT file.")

        if not text or len(text.strip()) < 40:
            raise ValueError("The uploaded document appears to be empty or could not be converted to readable text.")

        return text, file_type

parser_service = DocumentParserService()
